"""Reports (FR-30, FR-31): every type, scope, validation, preview limits and the Excel/PDF exports."""

import io
from datetime import UTC, date, datetime
from typing import Any

import pytest
from conftest import make_employee, make_org
from facetrack_common.constants import AttendanceStatus, ReportType
from facetrack_common.models import AttendanceDay
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import ValidationFailed
from app.core.security import DataScope
from app.domain.reports import export, service
from app.schemas.report import ReportFilters

ALL = DataScope(all_employees=True)


def _day(employee_id: int, day: int, status: AttendanceStatus, **minutes: int) -> AttendanceDay:
    check_in = (
        None
        if status in (AttendanceStatus.ABSENT, AttendanceStatus.ON_LEAVE)
        else datetime(2026, 9, day, 4, 30, tzinfo=UTC)
    )
    return AttendanceDay(
        employee_id=employee_id,
        work_date=date(2026, 9, day),
        status=status,
        check_in_at=check_in,
        check_out_at=datetime(2026, 9, day, 13, 0, tzinfo=UTC) if check_in else None,
        worked_minutes=minutes.get("worked", 510 if check_in else 0),
        late_minutes=minutes.get("late", 0),
        early_minutes=minutes.get("early", 0),
        overtime_minutes=minutes.get("overtime", 0),
        is_manual=minutes.get("manual", 0) == 1,
        finalized_at=datetime(2026, 9, day, 21, 0, tzinfo=UTC),
    )


@pytest.fixture
async def data(db: AsyncSession) -> dict[str, Any]:
    org = await make_org(db)
    created = datetime(2026, 8, 1, tzinfo=UTC)
    s1 = await make_employee(db, org, "S1", department="sales", created_at=created)
    s2 = await make_employee(db, org, "S2", department="sales", created_at=created)
    f1 = await make_employee(db, org, "F1", department="finance", created_at=created)
    db.add_all(
        [
            _day(s1.id, 1, AttendanceStatus.PRESENT),
            _day(s1.id, 2, AttendanceStatus.LATE, late=25),
            _day(s1.id, 3, AttendanceStatus.EARLY_EXIT, early=40, manual=1),
            _day(s2.id, 1, AttendanceStatus.ABSENT),
            _day(s2.id, 2, AttendanceStatus.PRESENT, overtime=45),
            _day(f1.id, 1, AttendanceStatus.ON_LEAVE),
            _day(f1.id, 2, AttendanceStatus.LATE, late=10),
        ]
    )
    await db.commit()
    return {"org": org, "s1": s1, "s2": s2, "f1": f1}


def _filters(**extra: Any) -> ReportFilters:
    return ReportFilters(date_from=date(2026, 9, 1), date_to=date(2026, 9, 30), **extra)


async def test_fr30_daily_report_rows_summary_and_chart(db: AsyncSession, data: dict[str, Any]) -> None:
    report = await service.build_report(db, ALL, ReportType.DAILY, _filters(), limit=500)
    assert report.total_rows == 7 and not report.truncated
    assert [r["employee_code"] for r in report.rows[:3]] == ["F1", "S1", "S2"]  # by date, then code
    summary = {s.label: s.value for s in report.summary}
    assert summary == {"Days": 7, "Attended": 5, "Late": 2, "Absent": 1, "On leave": 1, "Worked": 5 * 510}
    assert report.chart is not None and report.chart.data[0] == {
        "work_date": "2026-09-01",
        "attended": 1,
        "late": 0,
        "absent": 1,
    }


@pytest.mark.parametrize(
    ("report_type", "codes"),
    [
        (ReportType.LATE_ARRIVALS, ["S1", "F1"]),
        (ReportType.EARLY_EXITS, ["S1"]),
        (ReportType.ABSENTEE, ["S2"]),
        (ReportType.OVERTIME, ["S2"]),
    ],
)
async def test_fr30_exception_reports_select_the_right_days(
    db: AsyncSession, data: dict[str, Any], report_type: ReportType, codes: list[str]
) -> None:
    report = await service.build_report(db, ALL, report_type, _filters(), limit=500)
    assert sorted(r["employee_code"] for r in report.rows) == sorted(codes)


async def test_fr31_filters_by_department_employee_and_status(db: AsyncSession, data: dict[str, Any]) -> None:
    by_department = await service.build_report(
        db, ALL, ReportType.DAILY, _filters(department_id=data["org"]["finance"].id), limit=None
    )
    assert {r["employee_code"] for r in by_department.rows} == {"F1"}
    by_status = await service.build_report(
        db, ALL, ReportType.DAILY, _filters(status=AttendanceStatus.LATE), limit=None
    )
    assert {r["employee_code"] for r in by_status.rows} == {"S1", "F1"}
    history = await service.build_report(
        db, ALL, ReportType.EMPLOYEE_HISTORY, _filters(employee_id=data["s1"].id), limit=None
    )
    assert [r["work_date"] for r in history.rows] == ["2026-09-01", "2026-09-02", "2026-09-03"]


async def test_reports_respect_department_scope(db: AsyncSession, data: dict[str, Any]) -> None:
    manager_scope = DataScope(all_employees=False, department_ids=frozenset({data["org"]["sales"].id}))
    report = await service.build_report(db, manager_scope, ReportType.DAILY, _filters(), limit=None)
    assert {r["employee_code"] for r in report.rows} == {"S1", "S2"}
    self_scope = DataScope(all_employees=False, employee_id=data["f1"].id)
    mine = await service.build_report(db, self_scope, ReportType.ABSENTEE, _filters(), limit=None)
    assert mine.rows == []


async def test_department_summary_and_register(db: AsyncSession, data: dict[str, Any]) -> None:
    summary = await service.build_report(db, ALL, ReportType.DEPARTMENT_SUMMARY, _filters(), limit=None)
    sales = next(r for r in summary.rows if r["department"] == "Sales")
    assert sales["employees"] == 2 and sales["attended_days"] == 4 and sales["absent_days"] == 1
    assert sales["late_days"] == 1 and sales["overtime_minutes"] == 45
    register = await service.build_report(db, ALL, ReportType.MONTHLY_REGISTER, _filters(), limit=None)
    s1 = next(r for r in register.rows if r["employee_code"] == "S1")
    assert (s1["d01"], s1["d02"], s1["d03"], s1["total_P"]) == ("P", "L", "EE", 1)
    assert len([c for c in register.columns if c.kind == "letter"]) == 30


async def test_preview_is_capped_but_summary_covers_everything(
    db: AsyncSession, data: dict[str, Any]
) -> None:
    report = await service.build_report(db, ALL, ReportType.DAILY, _filters(), limit=2)
    assert len(report.rows) == 2 and report.total_rows == 7 and report.truncated
    assert {s.label: s.value for s in report.summary}["Days"] == 7


async def test_invalid_filters_are_rejected(db: AsyncSession) -> None:
    with pytest.raises(ValidationFailed):
        service.validate_filters(
            ReportType.DAILY, ReportFilters(date_from=date(2026, 9, 2), date_to=date(2026, 9, 1))
        )
    with pytest.raises(ValidationFailed):
        service.validate_filters(
            ReportType.DAILY, ReportFilters(date_from=date(2025, 1, 1), date_to=date(2026, 9, 1))
        )
    with pytest.raises(ValidationFailed):
        service.validate_filters(ReportType.EMPLOYEE_HISTORY, _filters())


async def test_fr31_excel_export_has_local_times_and_neutralised_text(
    db: AsyncSession, data: dict[str, Any]
) -> None:
    report = await service.build_report(db, ALL, ReportType.DAILY, _filters(), limit=None)
    report.rows[0]["employee_name"] = '=HYPERLINK("http://evil")'
    workbook = load_workbook(io.BytesIO(export.to_xlsx(report, "Department: Sales", "Hina HR")))
    sheet = workbook.worksheets[0]
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    header_index = next(i for i, row in enumerate(values) if row[0] == "Date")
    rows = values[header_index + 1 :]
    assert values[header_index][:7] == [
        "Date",
        "Code",
        "Employee",
        "Department",
        "Shift",
        "Check-in",
        "Check-out",
    ]
    s1 = next(row for row in rows if row[1] == "S1")
    assert s1[0] == "01-09-2026" and s1[5] == "09:30" and s1[6] == "18:00"  # 04:30/13:00 UTC in Asia/Karachi
    assert rows[0][2] == '\'=HYPERLINK("http://evil")'
    assert next(row for row in rows if row[1] == "F1")[5] is None  # on leave: no check-in
    assert workbook.sheetnames == ["Daily attendance", "Summary"]
    assert ["Generated by", "Hina HR"] in [row[:2] for row in values]


async def test_fr31_pdf_export_renders(db: AsyncSession, data: dict[str, Any]) -> None:
    report = await service.build_report(db, ALL, ReportType.DEPARTMENT_SUMMARY, _filters(), limit=None)
    pdf = export.to_pdf(report, "", "Hina HR")
    assert pdf.startswith(b"%PDF") and len(pdf) > 1_000


def _fill(cell: Any) -> str:
    return str(cell.fill.fgColor.rgb)[-6:].upper()


async def test_fr31_excel_export_colours_statuses_and_register_letters(
    db: AsyncSession, data: dict[str, Any]
) -> None:
    """Every sheet uses the portal's status colours: green present, yellow late, red absent (§13, Q62)."""
    daily = await service.build_report(db, ALL, ReportType.DAILY, _filters(), limit=None)
    workbook = load_workbook(io.BytesIO(export.to_xlsx(daily, "", "Hina HR")))
    sheet = workbook.worksheets[0]
    assert sheet.cell(1, 1).value == "FaceTrack — Daily attendance"
    header = next(row for row in sheet.iter_rows() if row[0].value == "Date")
    status_col = next(c.column for c in header if c.value == "Status")
    assert _fill(header[0]) == "1F4E8C" and header[0].font.color.rgb.endswith("FFFFFF")
    fills = {
        sheet.cell(r, status_col).value: _fill(sheet.cell(r, status_col))
        for r in range(header[0].row + 1, sheet.max_row + 1)
    }
    assert fills["Present"] == "C6EFCE" and fills["Late"] == "FFEB9C" and fills["Absent"] == "FFC7CE"
    assert fills["On leave"] == "E4DFEC"
    summary = workbook["Summary"]
    assert "Colour legend" in [row[0].value for row in summary.iter_rows()]

    register = await service.build_report(db, ALL, ReportType.MONTHLY_REGISTER, _filters(), limit=None)
    sheet = load_workbook(io.BytesIO(export.to_xlsx(register, "", "Hina HR"))).worksheets[0]
    letters = {
        c.value: _fill(c) for row in sheet.iter_rows() for c in row if c.value in ("P", "L", "A", "LV")
    }
    assert letters == {"P": "C6EFCE", "L": "FFEB9C", "A": "FFC7CE", "LV": "E4DFEC"}


def test_list_sheet_export_keeps_colours_numbers_and_neutralises_formulas() -> None:
    rows: list[list[export.Cell]] = [
        [("04-10-2026 17:03", None), ("Ehtisham Nasir", None), ("Recognised", "ok"), (0.85, None)],
        [("04-10-2026 17:04", None), ("=cmd|'/c calc'!A1", None), ("Unknown", "warn"), (None, None)],
    ]
    data = export.sheet_to_xlsx("Event log", ["Time", "Employee", "Status", "Score"], rows, "now", "Hina HR")
    workbook = load_workbook(io.BytesIO(data))
    sheet = workbook["Event log"]
    header = next(row for row in sheet.iter_rows() if row[0].value == "Time")
    first, second = header[0].row + 1, header[0].row + 2
    assert _fill(sheet.cell(first, 3)) == "C6EFCE" and _fill(sheet.cell(second, 3)) == "FFEB9C"
    assert sheet.cell(first, 4).value == 0.85 and sheet.cell(second, 4).value is None
    assert sheet.cell(second, 2).value == "'=cmd|'/c calc'!A1"
    summary = [[c.value for c in row] for row in workbook["Summary"].iter_rows()]
    assert ["Status: Recognised", 1] in summary and ["Status: Unknown", 1] in summary
