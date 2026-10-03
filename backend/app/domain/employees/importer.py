"""Bulk import (FR-11): employees from Excel/CSV, photos from a ZIP named by employee code.

Runs in Celery. Rows are upserted by `employee_code`; photos go through the normal enrollment
service (consent required, FR-9 checks). ZIP files are checked for path traversal, nested archives,
file count and decompression size (standards/14).
"""

import csv
import io
import logging
import re
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from pathlib import Path, PurePosixPath
from typing import Any

from facetrack_common.constants import EnrollmentSource
from facetrack_common.models import Department, Location, Shift, User
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.exceptions import DomainError
from app.domain.employees import enrollment
from app.domain.employees import service as employee_service
from app.repositories import employee_repo
from app.schemas.employee import EmployeeCreate, EmployeeUpdate

logger = logging.getLogger(__name__)

MAX_ROWS = 10_000
MAX_UNCOMPRESSED_BYTES = 2 * 1024**3
_COLUMNS = {
    "employee_code",
    "full_name",
    "department_code",
    "designation",
    "shift",
    "location",
    "email",
    "phone",
    "status",
    "consent_signed_at",
    "hr_external_id",
}
_PHOTO_NAME = re.compile(r"^(?P<code>[A-Za-z0-9][A-Za-z0-9._-]*?)(?:_\d+)?\.(?:jpe?g|png)$", re.IGNORECASE)


@dataclass
class ImportReport:
    created: int = 0
    updated: int = 0
    photos_accepted: int = 0
    photos_rejected: int = 0
    errors: list[dict[str, Any]] = field(default_factory=list)

    def error(self, where: str, message: str) -> None:
        if len(self.errors) < 200:
            self.errors.append({"where": where, "message": message})

    def as_dict(self) -> dict[str, Any]:
        return {
            "created": self.created,
            "updated": self.updated,
            "photos_accepted": self.photos_accepted,
            "photos_rejected": self.photos_rejected,
            "errors": self.errors,
            "error_count": len(self.errors),
        }


def read_rows(path: Path) -> list[dict[str, str]]:
    """Reads an .xlsx or .csv sheet into dicts keyed by normalised header names."""
    if path.suffix == ".xlsx":
        from openpyxl import load_workbook  # noqa: PLC0415

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        raw_rows = [["" if v is None else v for v in row] for row in sheet.iter_rows(values_only=True)]
        workbook.close()
    else:
        text = path.read_bytes().decode("utf-8-sig")
        raw_rows = list(csv.reader(io.StringIO(text)))
    if not raw_rows:
        return []
    headers = [str(h).strip().lower().replace(" ", "_") for h in raw_rows[0]]
    rows = []
    for raw in raw_rows[1 : MAX_ROWS + 1]:
        values = {h: _cell(v) for h, v in zip(headers, raw, strict=False) if h in _COLUMNS}
        if any(values.values()):
            rows.append(values)
    return rows


def _cell(value: object) -> str:
    if isinstance(value, datetime):
        return (value if value.tzinfo else value.replace(tzinfo=UTC)).isoformat()
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=UTC).isoformat()
    return str(value).strip()


def safe_zip_photos(path: Path, max_files: int) -> dict[str, list[tuple[str, bytes]]]:
    """Photos grouped by employee code; rejects traversal, nesting, too many files or bomb-like ZIPs."""
    photos: dict[str, list[tuple[str, bytes]]] = defaultdict(list)
    with zipfile.ZipFile(path) as archive:
        entries = [e for e in archive.infolist() if not e.is_dir()]
        if len(entries) > max_files:
            raise ValueError(f"the ZIP has more than {max_files} files")
        if sum(e.file_size for e in entries) > MAX_UNCOMPRESSED_BYTES:
            raise ValueError("the ZIP expands to more than 2 GB")
        for entry in entries:
            name = PurePosixPath(entry.filename)
            if name.is_absolute() or ".." in name.parts or "\\" in entry.filename:
                raise ValueError("the ZIP contains unsafe paths")
            if name.suffix.lower() in {".zip", ".rar", ".7z", ".gz", ".tar"}:
                raise ValueError("nested archives are not allowed")
            if name.name.startswith(".") or name.parts[0] == "__MACOSX":
                continue
            code = name.parts[0] if len(name.parts) > 1 else None
            if code is None:
                match = _PHOTO_NAME.match(name.name)
                code = match.group("code") if match else None
            if code is None or name.suffix.lower() not in {".jpg", ".jpeg", ".png"}:
                continue
            photos[code].append((name.name, archive.read(entry)))
    return photos


async def _lookup(db: AsyncSession, model: Any, column: Any, value: str) -> int | None:
    if not value:
        return None
    stmt = select(model.id).where(func.lower(column) == value.lower(), model.deleted_at.is_(None))
    return (await db.scalars(stmt)).first()


async def import_employees(
    db: AsyncSession, sheet_path: Path, zip_path: Path | None, actor: User
) -> dict[str, Any]:
    report = ImportReport()
    for number, row in enumerate(read_rows(sheet_path), start=2):
        where = f"row {number}"
        try:
            refs: dict[str, Any] = {
                "department_id": await _lookup(
                    db, Department, Department.code, row.get("department_code", "")
                ),
                "shift_id": await _lookup(db, Shift, Shift.name, row.get("shift", "")),
                "location_id": await _lookup(db, Location, Location.name, row.get("location", "")),
            }
            for key, label in (
                ("department_id", "department_code"),
                ("shift_id", "shift"),
                ("location_id", "location"),
            ):
                if row.get(label) and refs[key] is None:
                    raise ValueError(f"unknown {label} '{row[label]}'")
            data = {k: v for k, v in row.items() if v and k not in {"department_code", "shift", "location"}}
            data.update({k: v for k, v in refs.items() if v is not None})
            existing = await employee_repo.by_code(db, row.get("employee_code", ""))
            if existing is None:
                await employee_service.create_employee(db, EmployeeCreate.model_validate(data), actor)
                report.created += 1
            else:
                data.pop("employee_code", None)
                await employee_service.update_employee(
                    db, existing.id, EmployeeUpdate.model_validate(data), actor
                )
                report.updated += 1
        except ValidationError as exc:
            report.error(
                where, "; ".join(f"{'.'.join(map(str, e['loc']))}: {e['msg']}" for e in exc.errors())
            )
        except (DomainError, ValueError) as exc:
            report.error(where, getattr(exc, "message", None) or str(exc))

    if zip_path is not None:
        try:
            photos = safe_zip_photos(zip_path, get_settings().max_zip_files)
        except (ValueError, zipfile.BadZipFile) as exc:
            report.error("photos_zip", str(exc))
            photos = {}
        for code, files in photos.items():
            employee = await employee_repo.by_code(db, code)
            if employee is None:
                report.error(f"photos/{code}", "no employee with this code")
                continue
            try:
                results, _ = await enrollment.enroll_photos(
                    db,
                    employee.id,
                    [enrollment.IncomingPhoto(n, d) for n, d in files[:10]],
                    EnrollmentSource.UPLOAD,
                    actor,
                )
            except DomainError as exc:
                report.error(f"photos/{code}", exc.message)
                continue
            report.photos_accepted += sum(1 for r in results if r.accepted)
            report.photos_rejected += sum(1 for r in results if not r.accepted)
            for result in results:
                if not result.accepted:
                    report.error(f"photos/{code}/{result.filename}", result.rejection_reason or "rejected")
    return report.as_dict()
